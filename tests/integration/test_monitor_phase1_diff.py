from __future__ import annotations

import csv
from io import BytesIO, StringIO
from xml.etree import ElementTree as ET
from zipfile import ZIP_DEFLATED, ZipFile

from openpyxl import Workbook
from openpyxl.worksheet.table import Table
import pytest

from app.services.monitor_diff_service import MonitorDiffService, SvnMonitorSnapshotReader
from app.services.workbook_diff_service import DatasetLayout
from core.models import TreeEntry
from core.svn_history import BranchIdentity
from core.svn_provider import SVNProviderError


_SHEET_NS = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"


def _workbook_bytes(*, include_broken=False):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "main"
    sheet.append(["sheetName", "tbxName", "isExport"])
    sheet.append(["Role", "role", 1])
    sheet.append(["NotExported", "Missing", 0])
    if include_broken:
        sheet.append(["Broken", "Broken", 1])
    output = BytesIO()
    workbook.save(output)
    workbook.close()
    return output.getvalue()


def _broken_style_unbounded_manifest_bytes():
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "main"
    sheet.append(["sheetName", "tbxName", "isExport", "helper"])
    sheet.append(["Role", "role", 1, "outside table"])
    sheet.append(["Ghost", "GhostCsv", 1])
    sheet.add_table(Table(displayName="Manifest", ref="A1:C2"))
    source = BytesIO()
    workbook.save(source)
    workbook.close()

    target = BytesIO()
    with ZipFile(BytesIO(source.getvalue())) as input_archive:
        styles_root = ET.fromstring(input_archive.read("xl/styles.xml"))
        table_root = ET.fromstring(input_archive.read("xl/tables/table1.xml"))
        sheet_root = ET.fromstring(
            input_archive.read("xl/worksheets/sheet1.xml")
        )
        fills = styles_root.find(f"{{{_SHEET_NS}}}fills")
        assert fills is not None
        fill = ET.SubElement(fills, f"{{{_SHEET_NS}}}fill")
        gradient = ET.SubElement(fill, f"{{{_SHEET_NS}}}gradientFill")
        for color in ("FF000000", "FFFFFFFF"):
            stop = ET.SubElement(
                gradient,
                f"{{{_SHEET_NS}}}stop",
                {"position": "0"},
            )
            ET.SubElement(stop, f"{{{_SHEET_NS}}}color", {"rgb": color})
        fills.attrib["count"] = str(len(fills))
        table_root.attrib["ref"] = "1:2"
        auto_filter = table_root.find(f"{{{_SHEET_NS}}}autoFilter")
        if auto_filter is not None:
            auto_filter.attrib["ref"] = "1:2"
        dimension = sheet_root.find(f"{{{_SHEET_NS}}}dimension")
        if dimension is not None:
            sheet_root.remove(dimension)
        replacements = {
            "xl/styles.xml": ET.tostring(
                styles_root,
                encoding="utf-8",
                xml_declaration=True,
            ),
            "xl/tables/table1.xml": ET.tostring(
                table_root,
                encoding="utf-8",
                xml_declaration=True,
            ),
            "xl/worksheets/sheet1.xml": ET.tostring(
                sheet_root,
                encoding="utf-8",
                xml_declaration=True,
            ),
        }
        with ZipFile(target, "w", ZIP_DEFLATED) as output_archive:
            for item in input_archive.infolist():
                output_archive.writestr(
                    item,
                    replacements.get(
                        item.filename,
                        input_archive.read(item.filename),
                    ),
                )
    return target.getvalue()


def _csv_bytes(rows, *, hp_type="uint32", note="note"):
    output = StringIO(newline="")
    writer = csv.writer(output, lineterminator="\n")
    writer.writerows(
        [
            ["ID", "生命值", "策划备注"],
            ["Id", "Hp", "Note"],
            ["uint32", hp_type, "string"],
            ["All", "Client", "None"],
            ["meta", "meta", "meta"],
            ["meta", "meta", "meta"],
            ["meta", "meta", "meta"],
            *[[row[0], row[1], note] for row in rows],
        ]
    )
    return output.getvalue().encode("utf-8")


class History:
    def __init__(self, files):
        self.files = files
        self.reads = []

    def list_paths_at_revision(self, identity, revision):
        return [
            TreeEntry(path=path, kind="file")
            for path in self.files[revision]
        ]

    def read_path_bytes_at_revision(self, identity, path, revision):
        self.reads.append((path, revision))
        try:
            return self.files[revision][path]
        except KeyError as exc:
            raise SVNProviderError("SVN_PATH_NOT_FOUND", "missing") from exc


class FailingHistory(History):
    def __init__(self, files, code="SVN_TIMEOUT"):
        super().__init__(files)
        self.code = code

    def list_paths_at_revision(self, identity, revision):
        raise SVNProviderError(self.code, "private transport details")


class FailingReadHistory(History):
    def read_path_bytes_at_revision(self, identity, path, revision):
        raise SVNProviderError("SVN_NOT_REACHABLE", "private transport details")


def _layout():
    return DatasetLayout(
        csv_extension=".csv",
        filename_template="{tbxName}.csv",
        field_name_row=2,
        field_type_row=3,
        field_scope_row=4,
        data_start_row=8,
        primary_key_fields=("Id", "id"),
        manifest_sheet_name="main",
        manifest_sheet_field="sheetName",
        manifest_csv_name_field="tbxName",
        manifest_export_flag_field="isExport",
    )


def test_svn_snapshot_reader_reuses_m2_pairing_and_isolates_parse_failures():
    files = {
        100: {
            "Source/Table/Combat.xlsm": _workbook_bytes(),
            "Source/Table/Broken.xlsm": b"not-an-excel-file",
            "Source/TableCsv/Role.csv": _csv_bytes([("001", "01")], note="old"),
        },
        110: {
            "Source/Table/Combat.xlsm": _workbook_bytes(),
            "Source/Table/Broken.xlsm": b"still-not-an-excel-file",
            "Source/TableCsv/Role.csv": _csv_bytes(
                [("001", "1"), ("2", "80")],
                hp_type="uint64",
                note="changed but scope none",
            ),
        },
    }
    history = History(files)
    identity = BranchIdentity(
        canonical_url="https://svn.example/repo/branches/foo",
        repository_root="https://svn.example/repo",
        repository_uuid="20000000-0000-4000-8000-000000000001",
        repository_relative_path="branches/foo",
        bound_revision=110,
    )
    reader = SvnMonitorSnapshotReader(
        history,
        identity,
        _layout(),
        table_directory="Source/Table",
    )

    service = MonitorDiffService(reader)
    result = service.compare_revisions(100, 110)

    assert result.workbook_count == 2
    assert [(change.change_type.value, change.row_key, change.field_name) for change in result.changes] == [
        ("field_definition_modified", None, "Hp"),
        ("row_added", "2", None),
    ]
    assert result.changes[1].primary_key_field == "Id"
    assert result.changes[1].target.row_values == {"Id": "2", "Hp": "80"}
    assert [
        (field.field_name, field.display_name)
        for field in result.field_catalog[0].fields
    ] == [("Id", "ID"), ("Hp", "生命值")]

    history.reads.clear()
    selected_catalog = service.field_catalog_for_revisions(
        100,
        110,
        {("Combat.xlsm", "Role")},
    )
    assert selected_catalog == result.field_catalog
    assert not any("Broken.xlsm" in path for path, _ in history.reads)
    assert {
        (path, revision) for path, revision in history.reads
    } == {
        ("Source/Table/Combat.xlsm", 100),
        ("Source/TableCsv/Role.csv", 100),
        ("Source/Table/Combat.xlsm", 110),
        ("Source/TableCsv/Role.csv", 110),
    }
    assert [(error.workbook, error.stage.value) for error in result.errors] == [
        ("Broken.xlsm", "manifest_parse"),
    ]
    assert all(revision in {100, 110} for _, revision in history.reads)
    assert not any("Missing.csv" in path for path, _ in history.reads)


def test_broken_stylesheet_and_unbounded_manifest_stay_workbook_partial():
    broken = _broken_style_unbounded_manifest_bytes()
    files = {
        revision: {
            "Source/Table/Combat.xlsm": _workbook_bytes(),
            "Source/Table/BrokenStyles.xlsm": broken,
            "Source/TableCsv/Role.csv": _csv_bytes([("001", "100")]),
        }
        for revision in (100, 110)
    }
    history = History(files)
    identity = BranchIdentity(
        canonical_url="https://svn.example/repo/branches/synthetic",
        repository_root="https://svn.example/repo",
        repository_uuid="20000000-0000-4000-8000-000000000099",
        repository_relative_path="branches/synthetic",
        bound_revision=110,
    )
    reader = SvnMonitorSnapshotReader(
        history,
        identity,
        _layout(),
        table_directory="Source/Table",
    )

    result = MonitorDiffService(reader).compare_revisions(100, 110)

    assert result.workbook_count == 2
    assert result.reliable_workbook_count == 1
    assert result.changes == ()
    assert [
        (error.workbook, error.stage.value)
        for error in result.errors
    ] == [("BrokenStyles.xlsm", "manifest_parse")]


@pytest.mark.parametrize("provider_code", ("SVN_TIMEOUT", "SVN_NOT_REACHABLE"))
def test_svn_snapshot_reader_preserves_retryable_transport_error_classification(
    provider_code,
):
    identity = BranchIdentity(
        canonical_url="https://svn.example/repo/branches/foo",
        repository_root="https://svn.example/repo",
        repository_uuid="20000000-0000-4000-8000-000000000001",
        repository_relative_path="branches/foo",
        bound_revision=110,
    )
    reader = SvnMonitorSnapshotReader(
        FailingHistory({}, provider_code),
        identity,
        _layout(),
        table_directory="Source/Table",
    )

    snapshot = reader.load_snapshot(110)

    assert len(snapshot.errors) == 1
    assert snapshot.errors[0].code.value == "MONITOR_SVN_TIMEOUT"
    assert snapshot.errors[0].stage.value == "snapshot"
    assert snapshot.errors[0].retryable is True


def test_svn_snapshot_read_transport_failure_remains_retryable():
    identity = BranchIdentity(
        canonical_url="https://svn.example/repo/branches/foo",
        repository_root="https://svn.example/repo",
        repository_uuid="20000000-0000-4000-8000-000000000001",
        repository_relative_path="branches/foo",
        bound_revision=110,
    )
    history = FailingReadHistory(
        {
            110: {
                "Source/Table/Combat.xlsm": _workbook_bytes(),
                "Source/TableCsv/Role.csv": _csv_bytes([("1", "100")]),
            }
        }
    )
    reader = SvnMonitorSnapshotReader(
        history,
        identity,
        _layout(),
        table_directory="Source/Table",
    )

    snapshot = reader.load_snapshot(110)

    assert len(snapshot.errors) == 1
    assert snapshot.errors[0].code.value == "MONITOR_SVN_TIMEOUT"
    assert snapshot.errors[0].stage.value == "snapshot"
    assert snapshot.errors[0].retryable is True
    assert snapshot.errors[0].workbook == "Combat.xlsm"
